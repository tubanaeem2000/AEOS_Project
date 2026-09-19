"""
Text extraction for uploaded RAG documents.

This is new, isolated code - it does not modify or replace any existing
RAG file (retriever.py, rag_service.py, chunk_documents.py,
create_embeddings.py, create_vector_store.py). It only turns an uploaded
file's bytes into plain text, which then gets appended to the same
rag/documents/knowledge.txt the existing pipeline already reads.
"""
import csv
import io


class UnsupportedFileTypeError(Exception):
    pass


class FileExtractionError(Exception):
    pass


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".csv", ".xlsx"}


def extract_text(filename: str, file_bytes: bytes) -> str:
    """
    Returns the plain text content of an uploaded file.
    Raises UnsupportedFileTypeError or FileExtractionError on failure -
    callers should treat both as "could not process this file."
    """
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext or filename}'. "
            f"Supported types: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    try:
        if ext == ".txt":
            return _extract_txt(file_bytes)
        elif ext == ".csv":
            return _extract_csv(file_bytes)
        elif ext == ".pdf":
            return _extract_pdf(file_bytes)
        elif ext == ".docx":
            return _extract_docx(file_bytes)
        elif ext == ".xlsx":
            return _extract_xlsx(file_bytes)
    except (UnsupportedFileTypeError, FileExtractionError):
        raise
    except Exception as e:
        raise FileExtractionError(f"Failed to extract text from {filename}: {type(e).__name__}: {e}")

    raise UnsupportedFileTypeError(f"Unsupported file type '{ext}'")


def _extract_txt(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="replace")


def _extract_csv(file_bytes: bytes) -> str:
    text = file_bytes.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    lines = [", ".join(row) for row in reader if any(cell.strip() for cell in row)]
    return "\n".join(lines)


def _extract_pdf(file_bytes: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n\n".join(p.strip() for p in pages if p.strip())
    if not text.strip():
        raise FileExtractionError("No extractable text found in PDF (it may be a scanned/image-only PDF).")
    return text


def _extract_docx(file_bytes: bytes) -> str:
    import docx
    document = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


def _extract_xlsx(file_bytes: bytes) -> str:
    import openpyxl
    workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    lines = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(", ".join(cells))
    return "\n".join(lines)
